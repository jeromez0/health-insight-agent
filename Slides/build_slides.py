"""
Build a clean, consistently-formatted slides.pptx using python-pptx.
Replaces the pandoc-generated deck with a polished version.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE


# ---- design system ----

NAVY        = RGBColor(0x0A, 0x1F, 0x44)   # title bar / accents
INK         = RGBColor(0x1C, 0x1C, 0x1E)   # body text
SUBTLE      = RGBColor(0x6B, 0x6B, 0x70)   # secondary text
ACCENT      = RGBColor(0x00, 0x6E, 0xE6)   # links, highlights
SOFT_GRAY   = RGBColor(0xF2, 0xF2, 0xF7)   # alt row background
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
GREEN       = RGBColor(0x1F, 0x8A, 0x42)   # 100% / good
AMBER       = RGBColor(0xC9, 0x7B, 0x00)   # caution / variance

TITLE_FONT  = "Helvetica Neue"
BODY_FONT   = "Helvetica Neue"
MONO_FONT   = "Menlo"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


# ---- helpers ----

def set_slide_bg(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_title_bar(slide, text, subtext=None):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Inches(0.9))
    bar.line.fill.background()
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    tf = bar.text_frame
    tf.margin_left = Inches(0.6)
    tf.margin_right = Inches(0.6)
    tf.margin_top = Inches(0.18)
    tf.margin_bottom = Inches(0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = text
    r.font.name = TITLE_FONT
    r.font.size = Pt(26)
    r.font.bold = True
    r.font.color.rgb = WHITE
    if subtext:
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.LEFT
        r2 = p2.add_run()
        r2.text = subtext
        r2.font.name = BODY_FONT
        r2.font.size = Pt(13)
        r2.font.color.rgb = RGBColor(0xC4, 0xCC, 0xDB)


def add_textbox(slide, left, top, width, height, *, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    return tf


def add_run(p, text, *, size=18, bold=False, color=INK, font=BODY_FONT):
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return r


def add_para(tf, *, alignment=PP_ALIGN.LEFT, space_after=6, space_before=0, level=0, first=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = alignment
    p.space_after = Pt(space_after)
    p.space_before = Pt(space_before)
    p.level = level
    return p


def add_bullet(tf, text, *, level=0, size=18, bold=False, color=INK, first=False):
    p = add_para(tf, level=level, space_after=8, first=first)
    add_run(p, "•  ", size=size, bold=False, color=ACCENT)
    add_run(p, text, size=size, bold=bold, color=color)
    return p


def add_footer(slide, text):
    foot = slide.shapes.add_textbox(Inches(0.6), Inches(7.05), Inches(12.13), Inches(0.35))
    tf = foot.text_frame
    tf.margin_left = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = add_run(p, text, size=11, color=SUBTLE)


# ---- slides ----

def build_slide_1_title(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    set_slide_bg(s, WHITE)

    # left navy band
    band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.45), SLIDE_H)
    band.line.fill.background()
    band.fill.solid()
    band.fill.fore_color.rgb = NAVY

    # title block
    tf = add_textbox(s, Inches(1.0), Inches(2.4), Inches(11.5), Inches(3))
    p = add_para(tf, first=True)
    add_run(p, "Health Insight Agent", size=52, bold=True, color=NAVY)

    p = add_para(tf, space_before=20)
    add_run(p, "A Multi-Agent LLM Architecture for Personal Longitudinal Health Data Interpretation", size=22, color=INK)

    # author line
    tf2 = add_textbox(s, Inches(1.0), Inches(5.6), Inches(11.5), Inches(0.6))
    p = add_para(tf2, first=True)
    add_run(p, "Jerome Zhang", size=18, bold=True, color=INK)
    add_run(p, "   ·   MSAI High-Risk Project   ·   Spring 2026", size=18, color=SUBTLE)


def build_slide_2_problem(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(s, WHITE)
    add_title_bar(s, "The Problem")

    # Left column header
    tf = add_textbox(s, Inches(0.6), Inches(1.3), Inches(6.0), Inches(0.6))
    p = add_para(tf, first=True)
    add_run(p, "Personal health data is fragmented", size=22, bold=True, color=NAVY)

    tf = add_textbox(s, Inches(0.6), Inches(2.0), Inches(6.0), Inches(4.5))
    add_bullet(tf, "Wearables capture biometrics (HRV, sleep, body battery)", first=True)
    add_bullet(tf, "Food apps track nutrition (macros, micros)")
    add_bullet(tf, "Training journals log resistance work")
    p = add_para(tf, space_before=4)
    add_run(p, "    ", size=18)
    add_run(p, "…but no system unifies them into actionable insights", size=18, color=SUBTLE)

    # Right column header
    tf = add_textbox(s, Inches(7.0), Inches(1.3), Inches(6.0), Inches(0.6))
    p = add_para(tf, first=True)
    add_run(p, "Existing LLM approaches use single-shot prompting", size=22, bold=True, color=NAVY)

    tf = add_textbox(s, Inches(7.0), Inches(2.0), Inches(6.0), Inches(4.5))
    add_bullet(tf, "Dump everything into one prompt", first=True)
    add_bullet(tf, "Ask the LLM to figure it out")
    add_bullet(tf, "Generic, ungrounded outputs")
    add_bullet(tf, "Hidden cross-domain disagreements")

    add_footer(s, "Health Insight Agent  ·  Slide 2 / 8")


def build_slide_3_question(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(s, WHITE)
    add_title_bar(s, "Research Question")

    # quote-style emphasis box
    qbox = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(2.2), Inches(11.3), Inches(2.3))
    qbox.line.fill.background()
    qbox.fill.solid()
    qbox.fill.fore_color.rgb = SOFT_GRAY

    # left accent strip
    strip = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(2.2), Inches(0.12), Inches(2.3))
    strip.line.fill.background()
    strip.fill.solid()
    strip.fill.fore_color.rgb = ACCENT

    tf = qbox.text_frame
    tf.margin_left = Inches(0.4)
    tf.margin_right = Inches(0.4)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    add_run(p, "Does a multi-agent LLM architecture produce more grounded, domain-specific health insights than single-shot LLM prompting on identical longitudinal personal health data?", size=22, bold=True, color=NAVY)

    # approach
    tf = add_textbox(s, Inches(1.0), Inches(5.0), Inches(11.3), Inches(1.5))
    p = add_para(tf, first=True)
    add_run(p, "Approach", size=18, bold=True, color=SUBTLE)
    p = add_para(tf, space_before=6)
    add_run(p, "Build both. Compare side-by-side. Measure what differs.", size=18, color=INK)

    add_footer(s, "Health Insight Agent  ·  Slide 3 / 8")


def build_slide_4_architecture(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(s, WHITE)
    add_title_bar(s, "Method — Architecture")

    diagram = (
        "                    7-Day Personal Health Data\n"
        "              (Macros + Workouts + Garmin biometrics)\n"
        "                              │\n"
        "       ┌──────────────────────┼──────────────────────┐\n"
        "       ▼                      ▼                      ▼\n"
        "  ┌─────────┐           ┌──────────┐          ┌──────────┐\n"
        "  │Nutrition│           │ Training │          │ Recovery │       Stage 1\n"
        "  │  Agent  │           │  Agent   │          │  Agent   │      (parallel)\n"
        "  └─────┬───┘           └─────┬────┘          └─────┬────┘\n"
        "        │                     │                     │\n"
        "        └─────────────────────┼─────────────────────┘\n"
        "                              ▼\n"
        "                      ┌──────────────┐\n"
        "                      │ Orchestrator │                          Stage 2\n"
        "                      │    Agent     │\n"
        "                      └──────┬───────┘\n"
        "                             ▼\n"
        "                    3-5 Cross-Domain Insights"
    )

    box = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(1.2), Inches(8.3), Inches(5.0))
    box.line.color.rgb = RGBColor(0xE5, 0xE5, 0xEA)
    box.line.width = Pt(0.5)
    box.fill.solid()
    box.fill.fore_color.rgb = SOFT_GRAY
    tf = box.text_frame
    tf.margin_left = Inches(0.25)
    tf.margin_top = Inches(0.15)
    tf.margin_right = Inches(0.25)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    p.space_after = Pt(0)
    add_run(p, diagram, size=11, font=MONO_FONT, color=INK)

    # right column — comparison
    tf = add_textbox(s, Inches(9.4), Inches(1.4), Inches(3.6), Inches(5))
    p = add_para(tf, first=True)
    add_run(p, "Single-shot mode", size=16, bold=True, color=NAVY)
    p = add_para(tf, space_before=4)
    add_run(p, "1 LLM call. All data. Direct prompt → output.", size=14, color=INK)

    p = add_para(tf, space_before=20)
    add_run(p, "Multi-agent mode", size=16, bold=True, color=NAVY)
    p = add_para(tf, space_before=4)
    add_run(p, "4 LLM calls. 3 specialists in parallel + 1 orchestrator.", size=14, color=INK)

    add_footer(s, "Health Insight Agent  ·  Slide 4 / 8")


def build_slide_5_implementation(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(s, WHITE)
    add_title_bar(s, "Implementation")

    rows = [
        ("Platform", "Native iOS app  ·  SwiftUI  ·  iOS 17+"),
        ("LLM", "GPT-4o-mini via OpenAI Chat Completions API"),
        ("Networking", "Pure URLSession, no SDK"),
        ("Data", "7 days of self-tracked data (Apr 19–25, 2026)"),
        ("Domains", "Macros (food log) · Workouts (split + main lift) · Garmin (HRV, body battery, training readiness)"),
    ]

    top = Inches(1.4)
    row_h = Inches(0.7)
    for i, (label, value) in enumerate(rows):
        y = top + row_h * i
        # row alt-bg for readability
        if i % 2 == 0:
            r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), y, Inches(8.5), row_h)
            r.line.fill.background()
            r.fill.solid()
            r.fill.fore_color.rgb = SOFT_GRAY

        tf = add_textbox(s, Inches(0.8), y + Inches(0.12), Inches(2.0), Inches(0.5))
        p = add_para(tf, first=True)
        add_run(p, label, size=15, bold=True, color=SUBTLE)

        tf = add_textbox(s, Inches(2.7), y + Inches(0.12), Inches(6.3), Inches(0.5))
        p = add_para(tf, first=True)
        add_run(p, value, size=16, color=INK)

    # screenshot placeholder on right
    placeholder = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(9.4), Inches(1.4), Inches(3.4), Inches(5.4))
    placeholder.line.color.rgb = RGBColor(0xC7, 0xC7, 0xCC)
    placeholder.line.width = Pt(1)
    placeholder.line.dash_style = 7  # dashed
    placeholder.fill.solid()
    placeholder.fill.fore_color.rgb = WHITE
    tf = placeholder.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    add_run(p, "[App screenshot]", size=14, color=SUBTLE)
    p = add_para(tf, alignment=PP_ALIGN.CENTER, space_before=6)
    add_run(p, "Drop in the segmented control + insights card view", size=11, color=SUBTLE)

    add_footer(s, "Health Insight Agent  ·  Slide 5 / 8")


def build_slide_6_results(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(s, WHITE)
    add_title_bar(s, "Results", subtext="n=5 generations per mode  ·  GPT-4o-mini  ·  temp 0.4  ·  identical 7-day dataset")

    # table
    rows_data = [
        ("Metric", "Single-Shot", "Multi-Agent"),
        ("Grounding rate", "100%", "100%"),
        ("Cross-domain integration rate", "76%  (range 40–100%)", "100%  (no variance)"),
        ("Total tokens (mean)", "1352", "3603  (2.66×)"),
        ("Wall-clock latency (mean)", "9.80s", "13.90s  (1.42×)"),
        ("LLM calls / generation", "1", "4"),
    ]

    table_left = Inches(0.6)
    table_top = Inches(1.45)
    col_widths = [Inches(4.6), Inches(3.3), Inches(3.3)]
    row_h = Inches(0.5)

    rows = len(rows_data)
    cols = 3
    table_shape = s.shapes.add_table(rows, cols, table_left, table_top, sum(col_widths, Emu(0)), row_h * rows)
    table = table_shape.table
    for i, w in enumerate(col_widths):
        table.columns[i].width = w
    for r in range(rows):
        table.rows[r].height = row_h

    # styling per cell
    for r_idx, row in enumerate(rows_data):
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.margin_left = Inches(0.18)
            cell.margin_right = Inches(0.18)
            cell.margin_top = Inches(0.05)
            cell.margin_bottom = Inches(0.05)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if c_idx == 0 else PP_ALIGN.CENTER

            cell.fill.solid()
            if r_idx == 0:
                cell.fill.fore_color.rgb = NAVY
                add_run(p, val, size=14, bold=True, color=WHITE)
            else:
                cell.fill.fore_color.rgb = WHITE if r_idx % 2 == 1 else SOFT_GRAY
                bold = c_idx > 0 and (val == "100%" or "(no variance)" in val)
                color = GREEN if val == "100%" else INK
                add_run(p, val, size=14, bold=bold, color=color)

    # Headline finding
    tf = add_textbox(s, Inches(0.6), Inches(5.45), Inches(12.1), Inches(1.6))
    p = add_para(tf, first=True)
    add_run(p, "Headline finding", size=15, bold=True, color=SUBTLE)
    p = add_para(tf, space_before=6)
    add_run(p, "Both architectures hit 100% grounding (no hallucinated values). ", size=15, color=INK)
    add_run(p, "Multi-agent's win is consistency on cross-domain integration", size=15, bold=True, color=NAVY)
    add_run(p, " — it never dropped below 100%, while single-shot had one run at 40%. Cost: 2.66× tokens, 1.42× latency.", size=15, color=INK)

    add_footer(s, "Health Insight Agent  ·  Slide 6 / 8")


def build_slide_7_future(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(s, WHITE)
    add_title_bar(s, "Future Directions")

    items = [
        ("Scale up",            "Evaluate on public datasets (LifeSnaps, MyHeart Counts) across many users."),
        ("Clinical validation", "Pair LLM outputs with clinician annotation for correctness, not just grounding."),
        ("Adaptive routing",    "Dynamically choose which specialists to invoke based on user query."),
        ("On-device inference", "Swap GPT-4o-mini for Apple's Foundation Models framework — privacy-preserving deployment."),
        ("Persistent memory",   "Cross-session insight tracking (\"did you act on last week's recommendation?\")."),
    ]

    top = Inches(1.4)
    row_h = Inches(1.0)
    for i, (label, body) in enumerate(items):
        y = top + row_h * i
        # numbered chip
        chip = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.7), y + Inches(0.18), Inches(0.5), Inches(0.5))
        chip.line.fill.background()
        chip.fill.solid()
        chip.fill.fore_color.rgb = NAVY
        tf = chip.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_top = Inches(0.05); tf.margin_bottom = Inches(0); tf.margin_left = Inches(0); tf.margin_right = Inches(0)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        add_run(p, str(i+1), size=18, bold=True, color=WHITE)

        tf = add_textbox(s, Inches(1.4), y, Inches(11.4), row_h)
        p = add_para(tf, first=True)
        add_run(p, label, size=18, bold=True, color=NAVY)
        p = add_para(tf, space_before=2)
        add_run(p, body, size=14, color=INK)

    add_footer(s, "Health Insight Agent  ·  Slide 7 / 8")


def build_slide_8_thanks(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(s, WHITE)

    # left navy band
    band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.45), SLIDE_H)
    band.line.fill.background()
    band.fill.solid()
    band.fill.fore_color.rgb = NAVY

    tf = add_textbox(s, Inches(1.0), Inches(2.4), Inches(11.5), Inches(2))
    p = add_para(tf, first=True)
    add_run(p, "Thank You", size=60, bold=True, color=NAVY)

    tf = add_textbox(s, Inches(1.0), Inches(4.4), Inches(11.5), Inches(2))
    p = add_para(tf, first=True)
    add_run(p, "Jerome Zhang", size=20, bold=True, color=INK)
    add_run(p, "   ·   MSAI High-Risk Project   ·   Spring 2026", size=18, color=SUBTLE)

    p = add_para(tf, space_before=18)
    add_run(p, "Code & report:  ", size=16, color=INK)
    add_run(p, "https://github.com/jeromez0/health-insight-agent", size=16, color=ACCENT, font=MONO_FONT)

    p = add_para(tf, space_before=18)
    add_run(p, "Questions?", size=18, color=SUBTLE)


# ---- main ----

def main():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    build_slide_1_title(prs)
    build_slide_2_problem(prs)
    build_slide_3_question(prs)
    build_slide_4_architecture(prs)
    build_slide_5_implementation(prs)
    build_slide_6_results(prs)
    build_slide_7_future(prs)
    build_slide_8_thanks(prs)

    out = "slides.pptx"
    prs.save(out)
    print(f"wrote {out}  ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
